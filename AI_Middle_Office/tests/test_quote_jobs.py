import json
import re
import uuid
import asyncio
from datetime import datetime, timedelta, timezone
from io import BytesIO

import httpx
from openpyxl import Workbook

from app.core.config import settings
from app.core.database import SessionLocal
from app.core.security import get_password_hash
from app.models.client_inquiry import ClientInquiry
from app.models.cost_item import COST_STATUS_ACTIVE, CostItem
from app.models.quote_cost_evidence import QuoteCostEvidence
from app.models.quote_feedback import QuoteFeedback
from app.models.quote_history import QuoteHistory
from app.models.quote_job import QuoteJob, QuoteJobEvent
from app.models.quote_requirement_row import QuoteRequirementRow
from app.models.user import User, UserRole
from app.services import quote_job_runner
from app.services.quote_cost_matching import safe_enrich_quote_payload_with_cost_refs


def _login_headers(client):
    username = f"job_user_{uuid.uuid4().hex[:10]}"
    password = "secret123"
    response = client.post("/api/v1/auth/register", json={"username": username, "password": password})
    assert response.status_code == 200
    response = client.post("/api/v1/auth/login", data={"username": username, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _admin_headers(client):
    username = f"job_admin_{uuid.uuid4().hex[:10]}"
    password = "secret123"
    db = SessionLocal()
    try:
        db.add(User(username=username, hashed_password=get_password_hash(password), role="admin", quota=20))
        db.commit()
    finally:
        db.close()

    response = client.post("/api/v1/auth/login", data={"username": username, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _role_headers(client, roles: list[str], *, username_prefix: str = "job_role"):
    username = f"{username_prefix}_{uuid.uuid4().hex[:10]}"
    password = "secret123"
    legacy_role = "admin" if {"admin", "system_admin"} & set(roles) else "user"
    db = SessionLocal()
    try:
        user = User(username=username, hashed_password=get_password_hash(password), role=legacy_role, quota=20)
        db.add(user)
        db.flush()
        for role in roles:
            db.add(UserRole(user_id=user.id, role=role, created_by=None, note="test seed"))
        db.commit()
    finally:
        db.close()

    response = client.post("/api/v1/auth/login", data={"username": username, "password": password})
    assert response.status_code == 200
    return username, {"Authorization": f"Bearer {response.json()['access_token']}"}


def _response_data(response):
    return response.json()["data"]


def _quote_excel_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["工作内容", "工程量", "计量单位", "特征描述"])
    sheet.append(["拆除复合木地板", 20, "㎡", "不含清运"])
    sheet.append(["窗帘盒/灯槽拆除", 18, "m", "拆除至指定堆放点"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _set_flag(name: str, value):
    old_value = getattr(settings, name)
    object.__setattr__(settings, name, value)
    return old_value


def _seed_cost_item(db, *, item_name: str, unit: str = "m", price: float = 6.0) -> CostItem:
    item = CostItem(
        category="BIZ-2h 测试类",
        subcategory="成本前置",
        item_name=item_name,
        unit=unit,
        price=price,
        subcontract_composite_price=price,
        price_type="combined",
        status=COST_STATUS_ACTIVE,
        source="manual",
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def test_create_quote_job_returns_queued_status(client):
    headers = _login_headers(client)

    response = client.post(
        "/api/v1/quote/jobs",
        data={"message": "客厅地砖10平米"},
        headers=headers,
    )

    assert response.status_code == 202
    body = _response_data(response)
    assert body["job_id"]
    assert re.fullmatch(r"BJ-\d{8}-\d{6,}", body["job_number"])
    assert body["status"] == "queued"
    assert body["stage"] == "queued"
    assert body["trace_id"]
    assert body["request_summary"] == "客厅地砖10平米"
    assert body["result_item_count"] == 0
    assert body["events"][0]["status"] == "queued"

    status_response = client.get(f"/api/v1/quote/jobs/{body['job_id']}", headers=headers)
    assert status_response.status_code == 200
    detail = _response_data(status_response)
    assert detail["job_id"] == body["job_id"]
    assert detail["job_number"] == body["job_number"]
    assert detail["events"][0]["event_type"] == "queued"

    number_response = client.get(f"/api/v1/quote/jobs/{body['job_number']}", headers=headers)
    assert number_response.status_code == 200
    assert _response_data(number_response)["job_id"] == body["job_id"]

    search_response = client.get(
        "/api/v1/quote/jobs",
        params={"keyword": body["job_number"], "page_size": 20},
        headers=headers,
    )
    assert search_response.status_code == 200
    assert any(row["job_id"] == body["job_id"] for row in _response_data(search_response))

    db = SessionLocal()
    try:
        event = db.query(QuoteJobEvent).filter(QuoteJobEvent.quote_job_id == body["job_id"]).one()
        assert event.event_type == "queued"
        assert event.stage == "queued"
    finally:
        db.close()


def test_quote_operator_can_read_all_quote_jobs_without_mutating(client):
    owner_username, _ = _role_headers(client, ["staff"], username_prefix="job_owner")
    other_username, _ = _role_headers(client, ["staff"], username_prefix="job_other")
    _, operator_headers = _role_headers(client, ["quote_operator"], username_prefix="job_operator")

    owner_job_id = str(uuid.uuid4())
    other_job_id = str(uuid.uuid4())
    db = SessionLocal()
    try:
        db.add_all(
            [
                QuoteJob(
                    job_id=owner_job_id,
                    username=owner_username,
                    status="succeeded",
                    stage="completed",
                    message="operator visible owner job",
                    request_summary="operator visible owner job",
                ),
                QuoteJob(
                    job_id=other_job_id,
                    username=other_username,
                    status="queued",
                    stage="queued",
                    message="operator visible other job",
                    request_summary="operator visible other job",
                ),
            ]
        )
        db.commit()
    finally:
        db.close()

    list_response = client.get(
        "/api/v1/quote/jobs",
        headers=operator_headers,
        params={"keyword": "operator visible", "page_size": 50},
    )
    assert list_response.status_code == 200
    rows = _response_data(list_response)
    job_ids = {row["job_id"] for row in rows}
    assert {owner_job_id, other_job_id} <= job_ids

    detail_response = client.get(f"/api/v1/quote/jobs/{other_job_id}", headers=operator_headers)
    assert detail_response.status_code == 200
    assert _response_data(detail_response)["job_id"] == other_job_id

    review_response = client.get(f"/api/v1/quote/jobs/{other_job_id}/review-detail", headers=operator_headers)
    assert review_response.status_code == 200

    cancel_response = client.post(f"/api/v1/quote/jobs/{other_job_id}/cancel", headers=operator_headers)
    assert cancel_response.status_code == 404


def test_quote_job_review_detail_persists_confirmed_requirement_rows(client):
    headers = _login_headers(client)
    requirement_rows = [
        {
            "requirement_row_key": "装饰清单:10:0",
            "source_sheet": "装饰清单",
            "raw_row_index": 10,
            "item_name": "墙面白色腻子拆除",
            "spec": "按铲墙皮计价",
            "quantity": 12,
            "unit": "m2",
            "remark": "首层",
            "raw_text": "墙面白色腻子拆除 12 m2",
            "raw_cells": ["墙面白色腻子拆除", "12", "m2"],
        },
        {
            "requirement_row_key": "装饰清单:11:0",
            "source_sheet": "装饰清单",
            "raw_row_index": 11,
            "item_name": "墙面抹灰找平",
            "spec": "未含挂网",
            "quantity": 35,
            "unit": "m2",
            "remark": "二层",
            "raw_text": "墙面抹灰找平 35 m2",
            "raw_cells": ["墙面抹灰找平", "35", "m2"],
        },
    ]

    response = client.post(
        "/api/v1/quote/jobs",
        data={
            "message": "请根据确认清单报价",
            "requirement_rows_json": json.dumps(requirement_rows, ensure_ascii=False),
        },
        headers=headers,
    )

    assert response.status_code == 202
    job_id = _response_data(response)["job_id"]

    db = SessionLocal()
    try:
        stored_rows = (
            db.query(QuoteRequirementRow)
            .filter(QuoteRequirementRow.quote_job_id == job_id)
            .order_by(QuoteRequirementRow.sort_order.asc())
            .all()
        )
        assert [row.item_name for row in stored_rows] == ["墙面白色腻子拆除", "墙面抹灰找平"]
        assert stored_rows[0].source_sheet == "装饰清单"
        assert stored_rows[0].raw_row_index == 10

        job = db.query(QuoteJob).filter(QuoteJob.job_id == job_id).one()
        job.status = "succeeded"
        job.stage = "completed"
        job.result_json = json.dumps(
            {
                "project_details": [
                    {
                        "project_name": "墙面白色腻子拆除",
                        "quantity": 12,
                        "unit": "m2",
                        "unit_price": 12,
                        "total_price": 144,
                        "notes": "按铲墙皮计价",
                        "cost_reference": {"matched": False},
                    }
                ]
            },
            ensure_ascii=False,
        )
        db.commit()
    finally:
        db.close()

    detail_response = client.get(f"/api/v1/quote/jobs/{job_id}/review-detail", headers=headers)
    assert detail_response.status_code == 200
    detail = _response_data(detail_response)

    assert detail["summary"]["requirement_row_count"] == 2
    assert detail["summary"]["preview_row_count"] == 1
    assert detail["summary"]["matched_count"] == 1
    assert detail["summary"]["missing_count"] == 1
    assert detail["summary"]["no_cost_reference_count"] == 1
    assert detail["preview_rows"][0]["risk"]["label"] == "需复核"
    assert detail["preview_rows"][0]["checks"]["has_cost_reference"]["passed"] is False
    assert detail["missing_requirement_rows"][0]["requirement_row"]["item_name"] == "墙面抹灰找平"


def test_quote_job_review_detail_matches_requirement_row_key(client):
    headers = _login_headers(client)
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]
    job_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                result_json=json.dumps(
                    {
                        "project_details": [
                            {
                                "requirement_row_key": "Decor:10:0",
                                "project_name": "generic row A",
                                "quantity": 1,
                                "unit": "item",
                                "unit_price": 10,
                                "total_price": 10,
                                "notes": "matched by key",
                            },
                            {
                                "requirement_row_key": "Decor:11:0",
                                "project_name": "generic row B",
                                "quantity": 2,
                                "unit": "item",
                                "unit_price": 20,
                                "total_price": 40,
                                "notes": "matched by key",
                            },
                        ]
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.flush()
        db.add_all(
            [
                QuoteRequirementRow(
                    quote_job_id=job_id,
                    requirement_row_key="Decor:10:0",
                    source_sheet="Decor",
                    raw_row_index=10,
                    item_name="Requirement Alpha",
                    quantity=1,
                    unit="item",
                    sort_order=1,
                ),
                QuoteRequirementRow(
                    quote_job_id=job_id,
                    requirement_row_key="Decor:11:0",
                    source_sheet="Decor",
                    raw_row_index=11,
                    item_name="Requirement Beta",
                    quantity=2,
                    unit="item",
                    sort_order=2,
                ),
            ]
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}/review-detail", headers=headers)

    assert response.status_code == 200
    detail = _response_data(response)
    assert detail["summary"]["integrity_status"] == "complete"
    assert detail["summary"]["matched_count"] == 2
    assert detail["summary"]["missing_count"] == 0
    assert detail["missing_requirement_rows"] == []


def test_quote_job_review_detail_prefers_confirmed_manual_price(client):
    headers = _login_headers(client)
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]
    job_id = str(uuid.uuid4())
    quote_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                result_json=json.dumps(
                    {
                        "project_details": [
                            {
                                "project_name": "AI 未补价高风险项",
                                "quantity": 10,
                                "unit": "m2",
                                "unit_price": 0,
                                "total_price": 0,
                                "notes": "AI 未返回价格，人工已确认",
                                "cost_reference": {"matched": False},
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.flush()
        feedback = QuoteFeedback(
            quote_id=quote_id,
            quote_job_id=job_id,
            username=username,
            source="async_job",
            status="confirmed",
            pushed_to_dingtalk=True,
        )
        db.add(feedback)
        db.flush()
        db.add(
            QuoteCostEvidence(
                feedback_id=feedback.id,
                quote_id=quote_id,
                quote_job_id=job_id,
                username=username,
                source="async_job",
                status="confirmed",
                item_index=0,
                project_name="AI 未补价高风险项",
                quantity=10,
                unit="m2",
                ai_unit_price=0,
                ai_total_price=0,
                final_unit_price=58,
                final_total_price=580,
                line_total_price=580,
                line_total_source="manual_final",
                manual_modified=True,
            )
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}/review-detail", headers=headers)

    assert response.status_code == 200
    detail = _response_data(response)
    row = detail["preview_rows"][0]
    assert row["ai_unit_price"] == 0
    assert row["system_total_price"] == 0
    assert row["final_unit_price"] == 58
    assert row["final_total_price"] == 580
    assert row["display_unit_price"] == 58
    assert row["display_total_price"] == 580
    assert row["display_price_source"] == "final_confirmed"
    assert row["checks"]["ai_unit_price_positive"]["passed"] is True
    assert row["checks"]["ai_unit_price_positive"]["label"] == "最终确认单价大于 0"
    assert row["checks"]["system_total_positive"]["passed"] is True
    assert row["risk"]["type"] != "danger"
    assert detail["summary"]["high_risk_count"] == 0


def test_quote_job_review_detail_splits_manual_quantity_and_unit_price_risks(client):
    headers = _login_headers(client)
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]
    job_id = str(uuid.uuid4())
    quote_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                result_json=json.dumps(
                    {
                        "project_details": [
                            {
                                "project_name": "quantity changed item",
                                "quantity": 10,
                                "unit": "m2",
                                "unit_price": 100,
                                "total_price": 1000,
                                "notes": "quantity risk",
                                "cost_reference": {"matched": True, "reference_price": 100, "price_delta_rate": 0},
                            },
                            {
                                "project_name": "unit price changed item",
                                "quantity": 5,
                                "unit": "m",
                                "unit_price": 100,
                                "total_price": 500,
                                "notes": "price risk",
                                "cost_reference": {"matched": True, "reference_price": 100, "price_delta_rate": 0},
                            },
                        ]
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.flush()
        feedback = QuoteFeedback(
            quote_id=quote_id,
            quote_job_id=job_id,
            username=username,
            source="async_job",
            status="confirmed",
            pushed_to_dingtalk=True,
        )
        db.add(feedback)
        db.flush()
        db.add_all(
            [
                QuoteCostEvidence(
                    feedback_id=feedback.id,
                    quote_id=quote_id,
                    quote_job_id=job_id,
                    username=username,
                    source="async_job",
                    status="confirmed",
                    item_index=0,
                    project_name="quantity changed item",
                    quantity=14,
                    unit="m2",
                    ai_unit_price=100,
                    ai_total_price=1000,
                    final_unit_price=100,
                    final_total_price=1400,
                    line_total_price=1400,
                    line_total_source="manual_final",
                    manual_modified=True,
                ),
                QuoteCostEvidence(
                    feedback_id=feedback.id,
                    quote_id=quote_id,
                    quote_job_id=job_id,
                    username=username,
                    source="async_job",
                    status="confirmed",
                    item_index=1,
                    project_name="unit price changed item",
                    quantity=5,
                    unit="m",
                    ai_unit_price=100,
                    ai_total_price=500,
                    final_unit_price=140,
                    final_total_price=700,
                    line_total_price=700,
                    line_total_source="manual_final",
                    manual_modified=True,
                ),
            ]
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}/review-detail", headers=headers)

    assert response.status_code == 200
    rows = _response_data(response)["preview_rows"]
    quantity_row = rows[0]
    unit_price_row = rows[1]
    assert quantity_row["final_quantity"] == 14
    assert quantity_row["checks"]["manual_quantity_change_not_large"]["passed"] is False
    assert quantity_row["checks"]["manual_unit_price_change_not_large"]["passed"] is True
    assert unit_price_row["final_quantity"] == 5
    assert unit_price_row["checks"]["manual_quantity_change_not_large"]["passed"] is True
    assert unit_price_row["checks"]["manual_unit_price_change_not_large"]["passed"] is False


def test_confirm_push_rejects_incomplete_requirement_preview(client):
    headers = _login_headers(client)
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]
    job_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                result_json=json.dumps(
                    {
                        "project_details": [
                            {
                                "requirement_row_key": "Decor:10:0",
                                "project_name": "wall paint",
                                "quantity": 1,
                                "unit": "m2",
                                "unit_price": 10,
                                "total_price": 10,
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.flush()
        db.add_all(
            [
                QuoteRequirementRow(
                    quote_job_id=job_id,
                    requirement_row_key="Decor:10:0",
                    source_sheet="Decor",
                    raw_row_index=10,
                    item_name="wall paint",
                    quantity=1,
                    unit="m2",
                    sort_order=1,
                ),
                QuoteRequirementRow(
                    quote_job_id=job_id,
                    requirement_row_key="Decor:11:0",
                    source_sheet="Decor",
                    raw_row_index=11,
                    item_name="floor tile",
                    quantity=2,
                    unit="m2",
                    sort_order=2,
                ),
            ]
        )
        db.commit()
    finally:
        db.close()

    response = client.post(
        "/api/v1/confirm_push",
        json={
            "quote_job_id": job_id,
            "project_details": [
                {
                    "requirement_row_key": "Decor:10:0",
                    "project_name": "wall paint",
                    "unit_price": 10,
                    "total_price": 10,
                }
            ],
        },
        headers=headers,
    )

    assert response.status_code == 409
    assert "AI 预审不完整" in response.json()["detail"]


def test_confirm_push_rejects_unpriced_requirement_placeholders(client):
    headers = _login_headers(client)

    response = client.post(
        "/api/v1/confirm_push",
        json={
            "project_details": [
                {
                    "requirement_row_key": "Decor:12:0",
                    "project_name": "missing ai row",
                    "quantity": 2,
                    "unit": "m2",
                    "unit_price": 0,
                    "total_price": 0,
                    "requirement_placeholder": True,
                    "quote_source": "requirement_placeholder",
                }
            ],
        },
        headers=headers,
    )

    assert response.status_code == 409
    assert "占位行" in response.json()["detail"]


def test_cost_fallback_does_not_price_requirement_placeholder():
    suffix = uuid.uuid4().hex[:8]
    item_name = f"BIZ2l placeholder fallback {suffix}"
    db = SessionLocal()
    old_flag = _set_flag("feature_cost_db", True)
    try:
        _seed_cost_item(db, item_name=item_name, unit="m2", price=9.0)
        payload = {
            "project_details": [
                {
                    "project_name": item_name,
                    "quantity": 2,
                    "unit": "m2",
                    "unit_price": 0,
                    "total_price": 0,
                    "requirement_placeholder": True,
                    "quote_source": "requirement_placeholder",
                }
            ]
        }

        enriched = safe_enrich_quote_payload_with_cost_refs(db, payload, source_rows=[])
    finally:
        _set_flag("feature_cost_db", old_flag)
        db.close()

    row = enriched["project_details"][0]
    assert row["cost_reference"]["matched"] is True
    assert not row["cost_reference"].get("fallback_applied")
    assert row["unit_price"] == 0
    assert row["total_price"] == 0


def test_quote_job_rejects_invalid_requirement_rows_json(client):
    headers = _login_headers(client)

    response = client.post(
        "/api/v1/quote/jobs",
        data={"message": "报价", "requirement_rows_json": "{not-json"},
        headers=headers,
    )

    assert response.status_code == 400
    assert "requirement_rows_json" in response.json()["detail"]


def test_quote_job_review_detail_without_requirement_rows_does_not_report_extra_preview_rows(client):
    headers = _login_headers(client)
    job_id = str(uuid.uuid4())
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                message="手工报价",
                result_json=json.dumps(
                    {
                        "project_details": [
                            {
                                "project_name": "墙面抹灰找平",
                                "quantity": 35,
                                "unit": "m2",
                                "unit_price": 35,
                                "total_price": 1225,
                                "notes": "按抹灰找平计价",
                                "cost_reference": {"matched": False},
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
            )
        )
        db.commit()
    finally:
        db.close()

    detail_response = client.get(f"/api/v1/quote/jobs/{job_id}/review-detail", headers=headers)
    assert detail_response.status_code == 200
    detail = _response_data(detail_response)

    assert detail["summary"]["requirement_row_count"] == 0
    assert detail["summary"]["extra_count"] == 0
    assert detail["extra_preview_rows"] == []
    assert detail["reconciliation_rows"] == []
    assert detail["summary"]["no_cost_reference_count"] == 1


def test_quote_job_requires_auth(client):
    response = client.get("/api/v1/quote/jobs/missing")

    assert response.status_code == 401


def test_list_and_cancel_quote_job(client):
    headers = _login_headers(client)
    create_response = client.post("/api/v1/quote/jobs", data={"message": "卧室刷漆20平米"}, headers=headers)
    assert create_response.status_code == 202
    job_id = _response_data(create_response)["job_id"]

    list_response = client.get("/api/v1/quote/jobs?status=queued", headers=headers)
    assert list_response.status_code == 200
    assert any(item["job_id"] == job_id for item in list_response.json()["data"])

    cancel_response = client.post(f"/api/v1/quote/jobs/{job_id}/cancel", headers=headers)
    assert cancel_response.status_code == 200
    body = _response_data(cancel_response)
    assert body["status"] == "canceled"
    assert body["stage"] == "canceled"
    assert body["failure_stage"] == "canceled"
    assert body["duration_ms"] is not None
    assert body["error_message"] == "任务已取消"


def test_retry_canceled_quote_job_creates_new_job(client):
    headers = _login_headers(client)
    create_response = client.post("/api/v1/quote/jobs", data={"message": "厨房吊顶5平米"}, headers=headers)
    job_id = _response_data(create_response)["job_id"]
    cancel_response = client.post(f"/api/v1/quote/jobs/{job_id}/cancel", headers=headers)
    assert cancel_response.status_code == 200

    retry_response = client.post(f"/api/v1/quote/jobs/{job_id}/retry", headers=headers)
    assert retry_response.status_code == 202
    body = _response_data(retry_response)
    assert body["job_id"] != job_id
    assert body["status"] == "queued"
    assert body["events"][0]["source_job_id"] == job_id


def test_admin_can_list_all_jobs_and_mark_timeouts(client):
    user_headers = _login_headers(client)
    admin_headers = _admin_headers(client)
    create_response = client.post("/api/v1/quote/jobs", data={"message": "卫生间防水8平米"}, headers=user_headers)
    job_id = _response_data(create_response)["job_id"]

    db = SessionLocal()
    try:
        job = db.query(QuoteJob).filter(QuoteJob.job_id == job_id).first()
        job.status = "running"
        job.stage = "n8n"
        job.updated_at = datetime.now(timezone.utc) - timedelta(minutes=90)
        db.commit()
    finally:
        db.close()

    mark_response = client.post("/api/v1/admin/quote/jobs/mark_timeouts?timeout_minutes=30", headers=admin_headers)
    assert mark_response.status_code == 200
    body = mark_response.json()
    assert body["marked_count"] >= 1
    marked = next(item for item in body["data"] if item["job_id"] == job_id)
    assert marked["failure_stage"] == "timeout"
    assert marked["duration_ms"] is not None

    list_response = client.get("/api/v1/quote/jobs?status=timed_out", headers=admin_headers)
    assert list_response.status_code == 200
    assert any(item["job_id"] == job_id for item in list_response.json()["data"])


def test_admin_quote_job_list_includes_operational_context_and_filters(client):
    user_headers = _login_headers(client)
    admin_headers = _admin_headers(client)
    username = client.get("/api/v1/auth/me", headers=user_headers).json()["username"]
    job_id = str(uuid.uuid4())
    inquiry_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        user = db.query(User).filter(User.username == username).one()
        inquiry = ClientInquiry(
            inquiry_id=inquiry_id,
            source="微信",
            client_name="张三",
            client_phone="13800001111",
            inquiry_time=datetime.now(timezone.utc) - timedelta(minutes=20),
            first_response_time=datetime.now(timezone.utc),
            time_source="manual",
            responder_id=user.id,
            notes="需要局部翻新",
            first_quote_job_id=job_id,
        )
        job = QuoteJob(
            job_id=job_id,
            username=username,
            status="succeeded",
            stage="completed",
            message="厨房局部翻新",
            request_summary="厨房局部翻新",
            source_file_name="kitchen.png",
            result_total_amount=1200,
            result_item_count=3,
            client_inquiry_id=inquiry_id,
            duration_ms=1500,
            finished_at=datetime.now(timezone.utc),
        )
        history = QuoteHistory(
            username=username,
            quote_id=job_id,
            quote_job_id=job_id,
            confirmed_by=username,
            pushed_to_dingtalk=True,
            total_amount=1200,
            item_count=3,
            display_title="厨房局部翻新",
            project_summary="3 items",
            payload_json="{}",
        )
        db.add_all([inquiry, job, history])
        db.commit()
    finally:
        db.close()

    response = client.get(
        "/api/v1/quote/jobs?source=微信&keyword=13800001111&page_size=5",
        headers=admin_headers,
    )

    assert response.status_code == 200
    data = response.json()["data"]
    row = next(item for item in data if item["job_id"] == job_id)
    assert row["client_inquiry"]["client_name"] == "张三"
    assert row["client_inquiry"]["source"] == "微信"
    assert row["history"]["pushed_to_dingtalk"] is True
    assert row["history"]["total_amount"] == 1200


def test_quote_job_events_stream_replays_terminal_events(client):
    headers = _login_headers(client)
    login_response = client.get("/api/v1/auth/me", headers=headers)
    username = login_response.json()["username"]
    job_id = str(uuid.uuid4())
    events = [
        {"status": "queued", "message": "报价任务已进入队列", "trace_id": "trace-events", "stage": "queued"},
        {"status": "processing", "message": "异步报价任务已开始执行", "trace_id": "trace-events", "stage": "started"},
        {
            "status": "preview",
            "message": "AI 预审数据已就绪",
            "trace_id": "trace-events",
            "stage": "completed",
            "data": {"project_details": [{"project_name": "墙面刷新"}]},
        },
    ]

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="succeeded",
                stage="completed",
                trace_id="trace-events",
                events_json=json.dumps(events, ensure_ascii=False),
                result_json=json.dumps(events[-1]["data"], ensure_ascii=False),
                finished_at=datetime.now(timezone.utc),
            )
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}/events", headers=headers)

    assert response.status_code == 200
    streamed = [
        json.loads(line.removeprefix("data: "))
        for line in response.text.splitlines()
        if line.startswith("data: ")
    ]
    assert [event["status"] for event in streamed] == ["queued", "processing", "preview"]
    assert streamed[-1]["data"]["project_details"][0]["project_name"] == "墙面刷新"


def test_quote_job_detail_reads_structured_events_and_result_summary(client):
    headers = _login_headers(client)
    username = client.get("/api/v1/auth/me", headers=headers).json()["username"]
    job_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        job = QuoteJob(
            job_id=job_id,
            username=username,
            status="succeeded",
            stage="completed",
            message="living room renovation",
            file_name="plan.png",
            trace_id="trace-structured",
            result_json=json.dumps(
                {
                    "project_details": [
                        {"project_name": "wall paint", "total_price": 220},
                        {"project_name": "floor tile", "total_price": 110},
                    ]
                },
                ensure_ascii=False,
            ),
            request_summary="living room renovation",
            source_file_name="plan.png",
            result_total_amount=330,
            result_item_count=2,
            preview_project_names="wall paint, floor tile",
            duration_ms=1234,
            finished_at=datetime.now(timezone.utc),
        )
        db.add(job)
        db.flush()
        db.add(
            QuoteJobEvent(
                quote_job_id=job_id,
                event_index=1,
                event_type="preview",
                stage="completed",
                message="AI preview ready",
                trace_id="trace-structured",
                payload_json=json.dumps({"data": {"ok": True}}, ensure_ascii=False),
            )
        )
        db.add(
            QuoteFeedback(
                quote_id=job_id,
                quote_job_id=job_id,
                username=username,
                trace_id="trace-structured",
                status="rejected",
                rejected=True,
                rejection_reason="报价缺少找平厚度，请补充后重填。",
                reviewed_by=username,
                change_summary="Rejected: 报价缺少找平厚度，请补充后重填。",
                rejected_at=datetime.now(timezone.utc),
            )
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}", headers=headers)
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["request_summary"] == "living room renovation"
    assert data["source_file_name"] == "plan.png"
    assert data["result_total_amount"] == 330
    assert data["result_item_count"] == 2
    assert data["preview_project_names"] == ["wall paint", "floor tile"]
    assert data["duration_ms"] == 1234
    assert data["feedback"]["status"] == "rejected"
    assert data["feedback"]["rejected"] is True
    assert data["feedback"]["rejection_reason"] == "报价缺少找平厚度，请补充后重填。"
    assert data["feedback"]["reviewed_by"] == username
    assert data["events"][0]["event_type"] == "preview"
    assert data["events"][0]["payload"]["data"]["ok"] is True


def test_quote_job_events_hide_other_users_jobs(client):
    owner_headers = _login_headers(client)
    other_headers = _login_headers(client)
    owner = client.get("/api/v1/auth/me", headers=owner_headers).json()["username"]
    job_id = str(uuid.uuid4())

    db = SessionLocal()
    try:
        db.add(
            QuoteJob(
                job_id=job_id,
                username=owner,
                status="succeeded",
                stage="completed",
                trace_id="trace-private",
                events_json=json.dumps([{"status": "queued", "message": "private"}], ensure_ascii=False),
                finished_at=datetime.now(timezone.utc),
            )
        )
        db.commit()
    finally:
        db.close()

    response = client.get(f"/api/v1/quote/jobs/{job_id}/events", headers=other_headers)

    assert response.status_code == 200
    assert "报价任务不存在" in response.text


def test_quote_job_runner_records_runtime_duration(monkeypatch):
    username = f"runner_user_{uuid.uuid4().hex[:10]}"
    job_id = str(uuid.uuid4())

    async def fake_load_file_content(job, db):
        return None

    async def fake_quote_events(**kwargs):
        yield (
            "preview",
            "AI preview ready",
            {
                "stage": "completed",
                "data": {"project_details": [{"project_name": "墙面刷新", "total_price": 100}]},
            },
        )

    perf_values = iter([100.0, 102.345])
    monkeypatch.setattr(quote_job_runner.time, "perf_counter", lambda: next(perf_values))
    monkeypatch.setattr(quote_job_runner, "_load_job_file_content", fake_load_file_content)
    monkeypatch.setattr(quote_job_runner, "_iter_quote_events", fake_quote_events)
    monkeypatch.setattr(quote_job_runner, "safe_record_ai_preview", lambda *args, **kwargs: None)

    db = SessionLocal()
    try:
        db.add(User(username=username, hashed_password=get_password_hash("secret123"), role="user", quota=5))
        db.add(
            QuoteJob(
                job_id=job_id,
                username=username,
                status="queued",
                stage="queued",
                message="客厅墙面刷新",
                trace_id="trace-runtime-duration",
            )
        )
        db.commit()
    finally:
        db.close()

    asyncio.run(quote_job_runner.run_quote_job_async(job_id))

    db = SessionLocal()
    try:
        job = db.query(QuoteJob).filter(QuoteJob.job_id == job_id).one()
        user = db.query(User).filter(User.username == username).one()
        assert job.status == "succeeded"
        assert job.duration_ms == 2345
        assert job.result_item_count == 1
        assert user.quota == 4
    finally:
        db.close()


def test_quote_job_runner_normalizes_numbered_text_before_gateway(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(
            200,
            json={"project_details": [{"project_name": "拆除复合木地板", "unit_price": 7, "total_price": 245}]},
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message=(
                    "请生成报价明细，只包含以下三项：\n"
                    "1. 拆除复合木地板，35平方米\n"
                    "2. 拆除木脚线，42米\n"
                    "3. 拆砖墙（120厚砖墙），8平方米"
                ),
                file_content=None,
                mime_type=None,
                filename=None,
            )
        ]

    events = asyncio.run(collect_events())

    assert events[-1][0] == "preview"
    content = gateway_calls[0]["json_payload"]["text"]["content"]
    assert "\n" not in content
    assert "1." not in content
    assert "拆除复合木地板，35平方米；拆除木脚线，42米；拆砖墙（120厚砖墙），8平方米" in content


def test_quote_job_runner_sends_structured_requirement_rows_to_gateway(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(
            200,
            json={
                "project_details": [
                    {
                        "project_name": "wall paint",
                        "quantity": 1,
                        "unit": "m2",
                        "unit_price": 10,
                        "total_price": 10,
                    }
                ]
            },
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    requirement_rows = [
        QuoteRequirementRow(
            requirement_row_key="Decor:10:0",
            source_sheet="Decor",
            raw_row_index=10,
            item_name="wall paint",
            quantity=1,
            unit="m2",
            sort_order=1,
        ),
        QuoteRequirementRow(
            requirement_row_key="Decor:11:0",
            source_sheet="Decor",
            raw_row_index=11,
            item_name="floor tile",
            quantity=2,
            unit="m2",
            sort_order=2,
        ),
    ]

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message="quote confirmed rows",
                file_content=None,
                mime_type=None,
                filename=None,
                requirement_rows=requirement_rows,
            )
        ]

    events = asyncio.run(collect_events())

    assert gateway_calls
    content = gateway_calls[0]["json_payload"]["text"]["content"]
    assert "requirement_row_key=Decor:10:0" in content
    assert "requirement_row_key=Decor:11:0" in content
    assert "project_details" in content
    assert any(event[2].get("stage") == "requirement_guard" for event in events)
    payload = events[-1][2]["data"]
    details = payload["project_details"]
    assert [row["requirement_row_key"] for row in details] == ["Decor:10:0", "Decor:11:0"]
    assert details[0]["source_sheet"] == "Decor"
    assert details[0]["raw_row_index"] == 10
    assert details[1]["requirement_placeholder"] is True
    assert details[1]["unit_price"] == 0
    assert payload["requirement_merge_summary"]["required_count"] == 2
    assert payload["requirement_merge_summary"]["placeholder_count"] == 1
    assert payload["requirement_integrity"]["requirement_row_count"] == 2
    assert payload["requirement_integrity"]["matched_count"] == 2
    assert payload["requirement_integrity"]["missing_count"] == 0
    assert payload["requirement_integrity"]["status"] == "complete_with_placeholders"


def test_quote_job_runner_binds_requirement_rows_with_chinese_ai_fields(monkeypatch):
    async def fake_post_json_via_gateway(**kwargs):
        return httpx.Response(
            200,
            json={
                "project_details": [
                    {
                        "施工项目": "墙面腻子乳胶漆",
                        "工程量": 126.5,
                        "计量单位": "㎡",
                        "综合单价（元）": 22.94,
                        "合价": 2901.91,
                        "工艺备注": "AI returned with Chinese field names",
                    }
                ]
            },
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    requirement_rows = [
        QuoteRequirementRow(
            requirement_row_key="Decor:10:0",
            source_sheet="Decor",
            raw_row_index=10,
            item_name="墙面腻子乳胶漆",
            quantity=126.5,
            unit="㎡",
            sort_order=1,
        )
    ]

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message="quote confirmed rows",
                file_content=None,
                mime_type=None,
                filename=None,
                requirement_rows=requirement_rows,
            )
        ]

    events = asyncio.run(collect_events())

    payload = events[-1][2]["data"]
    details = payload["project_details"]
    assert payload["requirement_merge_summary"]["ai_returned_count"] == 1
    assert payload["requirement_merge_summary"]["placeholder_count"] == 0
    assert details[0]["requirement_row_key"] == "Decor:10:0"
    assert details[0]["project_name"] == "墙面腻子乳胶漆"
    assert details[0]["quantity"] == 126.5
    assert details[0]["unit"] == "㎡"
    assert details[0]["unit_price"] == 22.94
    assert details[0]["total_price"] == 2901.91
    assert details[0].get("requirement_placeholder") is not True
    assert payload["requirement_integrity"]["status"] == "complete"


def test_quote_job_runner_batches_requirement_rows_and_adds_placeholders(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        content = kwargs["json_payload"]["text"]["content"]
        keys = re.findall(r"requirement_row_key=([^|\s]+)", content)
        rows = [
            {
                "requirement_row_key": key,
                "project_name": f"quoted {key}",
                "quantity": 1,
                "unit": "m2",
                "unit_price": 10,
                "total_price": 10,
                "notes": "AI returned",
            }
            for key in keys
            if key not in {"Decor:3:0", "Decor:4:0"}
        ]
        return httpx.Response(200, json={"project_details": rows})

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    monkeypatch.setattr(quote_job_runner, "_quote_requirement_batch_size", lambda: 2)
    monkeypatch.setattr(quote_job_runner, "_quote_requirement_batch_threshold", lambda: 3)
    monkeypatch.setattr(quote_job_runner, "_quote_requirement_batch_retry_count", lambda: 1)
    monkeypatch.setattr(quote_job_runner, "_quote_job_heartbeat_interval_seconds", lambda: 0)
    requirement_rows = [
        QuoteRequirementRow(
            requirement_row_key=f"Decor:{index}:0",
            source_sheet="Decor",
            raw_row_index=index,
            item_name=f"item {index}",
            quantity=1,
            unit="m2",
            sort_order=index,
        )
        for index in range(1, 6)
    ]

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message="quote confirmed rows",
                file_content=None,
                mime_type=None,
                filename=None,
                requirement_rows=requirement_rows,
            )
        ]

    events = asyncio.run(collect_events())

    assert len(gateway_calls) == 4
    assert events[-1][0] == "preview"
    payload = events[-1][2]["data"]
    details = payload["project_details"]
    assert [row["requirement_row_key"] for row in details] == [f"Decor:{index}:0" for index in range(1, 6)]
    assert payload["requirement_integrity"]["status"] == "complete_with_placeholders"
    assert payload["requirement_integrity"]["missing_count"] == 0
    assert payload["requirement_integrity"]["placeholder_count"] == 2
    assert payload["requirement_batch_summary"]["batch_count"] == 3
    assert details[2]["requirement_placeholder"] is True
    assert details[2]["unit_price"] == 0
    assert details[3]["requirement_placeholder"] is True


def test_quote_job_runner_emits_heartbeat_while_waiting_for_n8n(monkeypatch):
    async def fake_post_json_via_gateway(**kwargs):
        await asyncio.sleep(0.05)
        return httpx.Response(
            200,
            json={"project_details": [{"project_name": "heartbeat quote", "unit_price": 7, "total_price": 70}]},
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    old_interval = _set_flag("quote_job_heartbeat_interval_seconds", 0.01)
    try:
        async def collect_events():
            return [
                event
                async for event in quote_job_runner._iter_quote_events(
                    username="runner",
                    message="heartbeat quote 10m",
                    file_content=None,
                    mime_type=None,
                    filename=None,
                )
            ]

        events = asyncio.run(collect_events())
    finally:
        _set_flag("quote_job_heartbeat_interval_seconds", old_interval)

    assert events[-1][0] == "preview"
    assert any(event[0] == "processing" and event[2].get("heartbeat") for event in events)


def test_quote_job_runner_uses_configured_n8n_timeout(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        raise httpx.ReadTimeout("slow n8n")

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    old_timeout = _set_flag("quote_n8n_timeout_seconds", 7)
    old_interval = _set_flag("quote_job_heartbeat_interval_seconds", 0)
    try:
        async def collect_events():
            return [
                event
                async for event in quote_job_runner._iter_quote_events(
                    username="runner",
                    message="timeout quote 10m",
                    file_content=None,
                    mime_type=None,
                    filename=None,
                )
            ]

        events = asyncio.run(collect_events())
    finally:
        _set_flag("quote_n8n_timeout_seconds", old_timeout)
        _set_flag("quote_job_heartbeat_interval_seconds", old_interval)

    assert gateway_calls[0]["timeout"] == 7
    assert events[-1][0] == "error"
    assert events[-1][2]["stage"] == "n8n"


def test_quote_job_runner_normalizes_plain_multiline_quote_items_before_gateway(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(
            200,
            json={"project_details": [{"project_name": "窗帘盒/灯槽拆除", "unit_price": 6, "total_price": 108}]},
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message="拆除复合木地板 20㎡\n拆除复合木地板 20㎡，拆除木脚线 30m\n窗帘盒/灯槽拆除 18m",
                file_content=None,
                mime_type=None,
                filename=None,
            )
        ]

    events = asyncio.run(collect_events())

    assert events[-1][0] == "preview"
    content = gateway_calls[0]["json_payload"]["text"]["content"]
    assert content == "拆除复合木地板 20㎡；拆除复合木地板 20㎡，拆除木脚线 30m；窗帘盒/灯槽拆除 18m"


def test_quote_job_runner_parses_excel_quote_sheet_before_gateway(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(
            200,
            json={"project_details": [{"project_name": "拆除复合木地板", "unit_price": 12, "total_price": 240}]},
        )

    async def fail_vision_call(*args, **kwargs):
        raise AssertionError("Excel quote sheets should not be sent to GLM-4V")

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)
    monkeypatch.setattr(quote_job_runner, "call_glm_vision_extract", fail_vision_call)

    async def collect_events():
        return [
            event
            async for event in quote_job_runner._iter_quote_events(
                username="runner",
                message="请根据需求单报价",
                file_content=_quote_excel_bytes(),
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                filename="quote.xlsx",
            )
        ]

    events = asyncio.run(collect_events())

    assert events[-1][0] == "preview"
    content = gateway_calls[0]["json_payload"]["text"]["content"]
    assert "从Excel需求单解析到的内容" in content
    assert "拆除复合木地板，规格/特征：不含清运，数量：20，单位：㎡" in content
    assert "窗帘盒/灯槽拆除，规格/特征：拆除至指定堆放点，数量：18，单位：m" in content


def test_quote_job_runner_attaches_biz2h_cost_context_before_gateway(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(
            200,
            json={"project_details": [{"project_name": "窗帘盒/灯槽拆除", "unit": "m", "quantity": 18, "unit_price": 6, "total_price": 108}]},
        )

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)

    db = SessionLocal()
    old_flag = _set_flag("feature_cost_db", True)
    try:
        _seed_cost_item(db, item_name="窗帘盒/灯槽拆除", unit="m", price=6.0)

        async def collect_events():
            return [
                event
                async for event in quote_job_runner._iter_quote_events(
                    username="runner",
                    message="请根据需求单报价",
                    file_content=_quote_excel_bytes(),
                    mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    filename="quote.xlsx",
                    db=db,
                )
            ]

        events = asyncio.run(collect_events())
    finally:
        _set_flag("feature_cost_db", old_flag)
        db.close()

    assert events[-1][0] == "preview"
    content = gateway_calls[0]["json_payload"]["text"]["content"]
    assert "[成本库底价强参考]" in content
    assert "需求项: 窗帘盒/灯槽拆除" in content
    assert "数量: 18" in content
    assert "匹配类型: fuzzy_item_name" in content
    assert "reference_unit_price: 6.00 元/m" in content


def test_quote_job_runner_falls_back_to_cost_context_when_n8n_returns_empty(monkeypatch):
    gateway_calls = []

    async def fake_post_json_via_gateway(**kwargs):
        gateway_calls.append(kwargs)
        return httpx.Response(200, content=b"")

    monkeypatch.setattr(quote_job_runner, "post_json_via_gateway", fake_post_json_via_gateway)

    suffix = uuid.uuid4().hex[:8]
    item_name = f"BIZ2h empty n8n fallback {suffix}"
    db = SessionLocal()
    old_flag = _set_flag("feature_cost_db", True)
    try:
        _seed_cost_item(db, item_name=item_name, unit="m", price=6.0)

        async def collect_events():
            return [
                event
                async for event in quote_job_runner._iter_quote_events(
                    username="runner",
                    message=f"{item_name} 18m",
                    file_content=None,
                    mime_type=None,
                    filename=None,
                    db=db,
                )
            ]

        events = asyncio.run(collect_events())
    finally:
        _set_flag("feature_cost_db", old_flag)
        db.close()

    assert gateway_calls
    assert events[-1][0] == "preview"
    assert events[-1][1].startswith("[Cost DB]")
    payload = events[-1][2]["data"]
    assert payload["cost_context_fallback_summary"]["reason"] == "n8n_empty_response"
    row = payload["project_details"][0]
    assert row["project_name"] == item_name
    assert row["quantity"] == "18"
    assert row["unit_price"] == 6
    assert row["total_price"] == 108
